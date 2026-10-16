"""Đối chiếu ba bên: Figma, repo GitHub và slide thuyết trình.

Chạy:
    python3 tools/sync_check.py --deck presentation/PRM323-Lab2-SE190533-slides.pptx \
        --figma-export <thư mục ảnh vừa xuất từ Figma> \
        --figma-facts process/sync/figma-facts.json \
        --out process/sync/bao-cao-doi-chieu.md

Cần: python-pptx, Pillow, numpy. Quy trình đầy đủ ở process/sync/README.md.
"""
import argparse, glob, io, json, os, re
import numpy as np
from PIL import Image
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def rd(p):
    return open(os.path.join(ROOT, p), encoding='utf-8').read()

def first_int(pattern, text):
    m = re.search(pattern, text)
    return int(m.group(1)) if m else None

def first_float(pattern, text):
    m = re.search(pattern, text)
    return float(m.group(1).replace(',', '.')) if m else None

# ---------- ảnh ----------
def vec(img, size=(48, 96)):
    im = img.convert('L')
    w, h = im.size
    if w > h:  # ảnh ngang (luồng, trang 04, 05, slide web)
        size = (96, max(8, int(96 * h / w)))
    else:
        size = (48, max(8, int(48 * h / w)))
    return np.asarray(im.resize(size, Image.LANCZOS), dtype=np.float32), im.size

def diff(a, b):
    if a.shape != b.shape:
        return 999.0
    return float(np.abs(a - b).mean())

def flat(im):
    # nền trong suốt coi như trắng, để ảnh xuất có góc trong suốt vẫn so được với ảnh nền trắng
    im = im.convert('RGBA')
    bg = Image.new('RGBA', im.size, (255, 255, 255, 255))
    return np.asarray(Image.alpha_composite(bg, im).convert('RGB'), dtype=np.int16)

def pixel_diff_img(a, b):
    d = np.abs(flat(a) - flat(b)).max(axis=2)
    return float((d > 8).mean() * 100)

def pixel_diff(p1, p2):
    a, b = Image.open(p1), Image.open(p2)
    if a.size != b.size:
        return None, f'{a.size} ≠ {b.size}'
    return pixel_diff_img(a, b), ''

def find_crop(img, big):
    """Tìm img như một phần cắt nguyên tỉ lệ của big. Trả về (sai khác, x, y)."""
    a = np.asarray(img.convert('L'), np.float32); b = np.asarray(big.convert('L'), np.float32)
    h, w = a.shape
    if b.shape[0] < h or b.shape[1] < w:
        return 999.0, 0, 0
    k = 8  # tìm thô trên ảnh thu nhỏ 8 lần, rồi dò lại quanh vị trí tốt nhất
    sa, sb = a[::k, ::k], b[::k, ::k]
    win = np.lib.stride_tricks.sliding_window_view(sb, sa.shape)
    d = np.abs(win - sa).mean(axis=(2, 3))
    y0, x0 = np.unravel_index(d.argmin(), d.shape)
    best = (999.0, 0, 0)
    for y in range(max(0, y0 * k - k), min(b.shape[0] - h, y0 * k + k) + 1):
        for x in range(max(0, x0 * k - k), min(b.shape[1] - w, x0 * k + k) + 1):
            v = float(np.abs(b[y:y + h, x:x + w] - a).mean())
            if v < best[0]:
                best = (v, x, y)
    return best

def level(score):
    return 'khớp' if score < 3 else ('gần khớp' if score < 10 else 'LỆCH')

# ---------- số liệu ----------
def contrast(h1, h2):
    def lum(h):
        c = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
        c = [x / 12.92 if x <= 0.04045 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
    a, b = sorted([lum(h1), lum(h2)], reverse=True)
    return round((a + 0.05) / (b + 0.05), 2)

def repo_numbers():
    dl = rd('process/decision-log.md')
    crs = re.findall(r'CR-(\d+)', rd('ai/ai-design-log.md') + rd('design/final-ui-review.md'))
    pl = rd('design/prototype-links.md')
    idx = rd('assets/figma/INDEX.md')
    dd = rd('design/design-decisions.md')
    spec = rd('design/screen-spec.md')
    des = rd('design/DESIGN.md')
    return {
        'man': len(set(re.findall(r'^#{2,3} (SC-\d+)', spec, re.M))),
        'cr': max(int(x) for x in crs),
        'quyet_dinh': len(set(re.findall(r'^\| D-(\d+) ', dl, re.M))),
        'frame_360': first_int(r'Tổng: (\d+) frame 360', idx),
        'frame_dien_thoai': first_int(r'trên (\d+) frame điện thoại', dd),
        'component_set': len(set(re.findall(r'\| (CMP-\d+) \|', des))),
        'bien_mau': len(set(re.findall(r'`(color/[a-z-]+)`', des))),
        'proto_frame': first_int(r'Frame trên trang 06 \| (\d+)', pl),
        'diem_bat_dau': first_int(r'Điểm bắt đầu[^|]*\| (\d+)', pl),
        'tuong_tac': first_int(r'\| Tương tác \| (\d+)', pl),
        'vung_bam': first_int(r'(\d+) vùng có hành động chạm', dd),
        'tuong_phan_chu': first_float(r'đạt ít nhất (\d+,\d+):1', dd),
        'tuong_phan_vien': first_float(r'`border-control` (\d+,\d+):1', dd),
    }

def figma_numbers(f):
    m = f['mau_light']
    return {
        'man': f['trang_03']['so_man_SC'],
        'frame_360': f['trang_03']['kich_thuoc']['360x800'],
        'frame_dien_thoai': f['trang_03']['kich_thuoc']['360x800'] + f['trang_03']['kich_thuoc']['412x915'],
        'component_set': f['trang_05']['component_set'],
        'text_style': f['text_style'],
        'do_noi': f['effect_style'],
        'bien_mau': f['bien']['Color · DESIGN.md v2'],
        'proto_frame': f['trang_06']['frame'],
        'diem_bat_dau': f['trang_06']['diem_bat_dau'],
        'tuong_tac': f['trang_06']['tuong_tac_tho'] - f['trang_06']['tuong_tac_bi_che'],
        'tuong_phan_chu': contrast(m['color/warning'], m['color/canvas']),
        'tuong_phan_vien': contrast(m['color/border-control'], m['color/surface']),
    }

CLAIMS = [  # (khoá, slide, nhãn, regex trên chữ của slide)
    ('man', 3, 'Số màn', r'(\d+) màn'),
    ('cr', 6, 'Phát hiện có mã CR', r'(\d+) phát hiện'),
    ('quyet_dinh', 7, 'Quyết định D-nn', r'(\d+) quyết định'),
    ('frame_360', 8, 'Frame 360 × 800 trang 03', r'(\d+) màn và biến thể 360'),
    ('frame_dien_thoai', 9, 'Frame điện thoại đổi theo token', r'cả (\d+) frame'),
    ('text_style', 9, 'Text style', r'(\d+) text style'),
    ('do_noi', 9, 'Độ nổi (effect style)', r'(\d+) độ nổi'),
    ('component_set', 9, 'Component set', r'(\d+) component set'),
    ('bien_mau', 9, 'Biến màu', r'(\d+) biến màu'),
    ('vung_bam', 10, 'Vùng bấm 48 dp', r'\((\d+) vùng\)'),
    ('tuong_phan_chu', 10, 'Tương phản chữ thấp nhất', r'(\d+,\d+):1 Tương phản chữ'),
    ('tuong_phan_vien', 10, 'Tương phản viền điều khiển', r'(\d+,\d+):1 Viền'),
    ('proto_frame', 11, 'Frame prototype trang 06', r'(\d+) frame ·'),
    ('diem_bat_dau', 11, 'Điểm bắt đầu flow', r'(\d+) điểm bắt đầu'),
    ('tuong_tac', 11, 'Tương tác prototype', r'(\d+) tương tác'),
]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--deck', required=True)
    ap.add_argument('--figma-export', required=True)
    ap.add_argument('--figma-facts', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    deck = Presentation(a.deck)
    facts = json.load(open(a.figma_facts, encoding='utf-8'))
    out = ['# Báo cáo đối chiếu Figma · GitHub · slide', '',
           f'Slide: `{os.path.relpath(a.deck, ROOT)}` · Figma: {facts["_nguon"]}', '']

    # 1. Figma (ảnh vừa xuất) ↔ repo assets/figma
    out += ['## 1. Ảnh Figma vừa xuất ↔ `assets/figma`', '',
            'So từng pixel; “% điểm khác” là tỉ lệ điểm ảnh lệch quá 8/255 ở một kênh màu.', '',
            '| Ảnh | Kết quả |', '|---|---|']
    fresh = sorted(glob.glob(os.path.join(a.figma_export, '*.png')))
    bad1 = 0
    for p in fresh:
        name = os.path.basename(p)
        rp = os.path.join(ROOT, 'assets/figma', name)
        if not os.path.exists(rp):
            out.append(f'| {name} | LỆCH: repo không có ảnh này |'); bad1 += 1; continue
        pct, note = pixel_diff(p, rp)
        if pct is None:
            out.append(f'| {name} | LỆCH: kích thước {note} |'); bad1 += 1
        elif pct > 0.5:
            out.append(f'| {name} | LỆCH: {pct:.2f}% điểm khác |'); bad1 += 1
    repo_imgs = {os.path.basename(p) for p in glob.glob(os.path.join(ROOT, 'assets/figma/*.png'))}
    for name in sorted(repo_imgs - {os.path.basename(p) for p in fresh}):
        out.append(f'| {name} | Chưa xuất lại từ Figma lần này |')
    if out[-1] == '|---|---|':  # không có dòng nào: bỏ bảng rỗng
        out = out[:-3]
    out += ['', f'**{len(fresh) - bad1}/{len(fresh)} ảnh khớp** (lệch dưới 0,5% điểm ảnh).', '']

    # 2. slide ↔ repo assets
    lib, big = {}, {}
    for p in glob.glob(os.path.join(ROOT, 'assets/figma/*.png')) + glob.glob(os.path.join(ROOT, 'assets/stitch/*.png')):
        im = Image.open(p)
        lib[os.path.relpath(p, ROOT)] = vec(im)
        big[os.path.relpath(p, ROOT)] = im
    out += ['## 2. Ảnh trên slide ↔ ảnh trong repo', '',
            'Mỗi ảnh chụp màn hình trên slide được so với mọi ảnh trong `assets/figma` và `assets/stitch` (thu nhỏ, thang xám, sai khác trung bình trên 255). Ảnh cùng kích thước với ảnh repo thì so từng điểm ảnh như mục 1. Ảnh khác kích thước thì dưới 3 là khớp, 3–10 là gần khớp (thường do slide nén ảnh), trên 10 là lệch. Ảnh không khớp nguyên tấm được dò như một phần cắt của ảnh trang lớn hơn.', '',
            '| Slide | Ảnh trên slide | Ảnh repo giống nhất | Sai khác | Kết quả |', '|---|---|---|---|---|']
    bad2 = 0; n2 = 0; icons = 0
    for si, sl in enumerate(deck.slides, 1):
        for sh in sl.shapes:
            if sh.shape_type != 13:
                continue
            if max(sh.width, sh.height) < 40 * 12700:  # biểu tượng Lucide, dưới 40 pt trên slide
                icons += 1; continue
            img = Image.open(io.BytesIO(sh.image.blob))
            v, _ = vec(img)
            best = min(((diff(v, lv[0]), k) for k, lv in lib.items()), default=(999, '-'))
            where = f'`{best[1]}`'
            if level(best[0]) == 'LỆCH':  # có thể là phần cắt từ một trang lớn
                for k, im in big.items():
                    c = find_crop(img, im)
                    if c[0] < best[0]:
                        best = (c[0], k); where = f'`{k}`, phần cắt tại x={c[1]}, y={c[2]}'
            n2 += 1
            lv = level(best[0]); score = f'{best[0]:.1f}'
            if best[1] in big and big[best[1]].size == img.size:
                # cùng kích thước: so từng điểm ảnh để bắt cả thay đổi nhỏ như một dấu gạch
                pct = pixel_diff_img(img, big[best[1]])
                lv = 'khớp' if pct <= 0.5 else 'LỆCH'; score = f'{pct:.2f}% điểm'
            bad2 += lv == 'LỆCH'
            descr = (sh._element.xpath('./p:nvPicPr/p:cNvPr/@descr') or [sh.name])[0]
            out.append(f'| {si} | {descr} | {where} | {score} | {lv} |')
    out += ['', f'**{n2 - bad2}/{n2} ảnh trên slide khớp hoặc gần khớp với repo.** Bỏ qua {icons} biểu tượng nhỏ hơn 40 pt (vẽ từ Lucide, không có trong Figma).', '']

    # 3. số liệu
    rn, fn = repo_numbers(), figma_numbers(facts)
    texts = {i: ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for i, s in enumerate(deck.slides, 1)}
    out += ['## 3. Số liệu trên slide ↔ repo ↔ Figma', '',
            '| Số liệu | Slide | Trên slide | Repo | Figma | Kết quả |', '|---|---|---|---|---|---|']
    bad3 = 0
    for key, si, label, rx in CLAIMS:
        m = re.search(rx, texts[si], re.S)
        sv = m.group(1) if m else None
        svn = float(sv.replace(',', '.')) if sv else None
        r, f = rn.get(key), fn.get(key)
        vals = [x for x in (svn, r, f) if x is not None]
        ok = svn is not None and all(abs(x - svn) < 0.006 for x in vals)
        bad3 += not ok
        fmt = lambda x: '—' if x is None else (f'{x:g}'.replace('.', ',') if isinstance(x, float) else str(x))
        out.append(f'| {label} | {si} | {fmt(svn)} | {fmt(r)} | {fmt(f)} | {"khớp" if ok else "LỆCH"} |')
    out += ['', f'**{len(CLAIMS) - bad3}/{len(CLAIMS)} số liệu khớp.** “—” là bên đó không có số này để so.', '',
            '“Tương tác prototype” ở cột Figma là số tương tác thô trừ các tương tác bị hotspot đè (xem `figma-facts.json`).', '']
    open(a.out, 'w', encoding='utf-8').write('\n'.join(out))
    print(f'anh_figma_lech={bad1} anh_slide_lech={bad2} so_lieu_lech={bad3}')

if __name__ == '__main__':
    main()
